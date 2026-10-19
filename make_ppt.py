# -*- coding: utf-8 -*-
"""生成 5 分钟结题答辩 PPT（16:9，17 页）。

叙事：
  封面 -> 问题 -> 评估基座 -> 核心思路/流程 -> 物理粗提取 ->
  框提示/SAM 精修 -> LoRA -> 实验设置 -> 主结果表 -> 主结果图/配对 ->
  消融 -> 跨数据集/误差 -> 单图多方法 -> 软件原型 ->
  方向先验/S-EO -> 结论/边界 -> 致谢

口径纪律：
  - 主结果 = 无泄漏 68/16/21，LoRA 报 3 seed 均值；基线 3 seed
  - 只说“数值上高于”，不宣称 SOTA/显著超越
  - 方向模块写“预注册试点未达到阳性判据”，不写“方向可学/方向先验有效”
"""
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
DARK = RGBColor(0x1F, 0x3B, 0x63)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)
GRAY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xA6, 0x3A, 0x2E)
LIGHT = RGBColor(0xDE, 0xEA, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

PROJECT = "基于方向提示与 SAM 的遥感影像阴影检测方法研究"
FOOTER = "ShadowHunter · 结题答辩 · 2026 年 10 月"


def set_font(run, size=16, bold=False, color=None, name="Microsoft YaHei"):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = name
    if color is not None:
        f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def title_bar(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12.15), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_font(r, size=25, bold=True, color=DARK)
    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = sub
        set_font(r2, size=12, color=GRAY)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_footer(slide, text=FOOTER):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.12), Inches(12.15), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_font(r, size=9, color=GRAY)


def add_bullets(tf, items, size=16, color=DARK):
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = it
        set_font(r, size=size, color=RED if it[:1] in ("!", "！") else color)
        p.space_after = Pt(7)


def content_slide(title, bullets, notes="", sub=None, size=16):
    s = prs.slides.add_slide(BLANK)
    title_bar(s, title, sub)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(12.13), Inches(5.35))
    tf = tb.text_frame
    tf.word_wrap = True
    add_bullets(tf, bullets, size=size)
    add_notes(s, notes)
    add_footer(s)
    return s


def table_slide(title, rows, widths=None, notes="", sub=None, note=None,
                row_h=0.56, top=1.5, best_row=1, size=13):
    s = prs.slides.add_slide(BLANK)
    title_bar(s, title, sub)
    n_r, n_c = len(rows), len(rows[0])
    shape = s.shapes.add_table(n_r, n_c, Inches(0.6), Inches(top),
                               Inches(12.13), Inches(row_h * n_r))
    tbl = shape.table
    if widths:
        total = sum(widths)
        for j, wj in enumerate(widths):
            tbl.columns[j].width = Inches(12.13 * wj / total)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(val)
            pr = cell.text_frame.paragraphs[0]
            pr.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            is_head = (i == 0)
            is_best = (best_row is not None and i == best_row)
            set_font(pr.runs[0], size=size, bold=(is_head or is_best),
                     color=WHITE if is_head else (ACCENT if is_best else DARK))
            if is_head:
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT
            elif is_best:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT
    if note:
        tb = s.shapes.add_textbox(Inches(0.6), Inches(top + row_h * n_r + 0.12),
                                  Inches(12.13), Inches(max(0.5, 7.0 - top - row_h * n_r)))
        tf = tb.text_frame
        tf.word_wrap = True
        add_bullets(tf, note, size=10)
    add_notes(s, notes)
    add_footer(s)
    return s


def image_slide(title, image, bullets=None, notes="", sub=None, mode="left",
                caption=None, image_h=5.2):
    s = prs.slides.add_slide(BLANK)
    title_bar(s, title, sub)
    img_path = BASE / image
    if mode == "full":
        pic = s.shapes.add_picture(str(img_path), Inches(0.6), Inches(1.4), width=Inches(12.13))
        if pic.height > Inches(5.1):
            scale = Inches(5.1) / pic.height
            pic.height = Inches(5.1)
            pic.width = int(pic.width * scale)
    else:
        pic = s.shapes.add_picture(str(img_path), Inches(0.6), Inches(1.45), height=Inches(image_h))
        if pic.width > Inches(7.0):
            scale = Inches(7.0) / pic.width
            pic.width = Inches(7.0)
            pic.height = int(pic.height * scale)
        tb = s.shapes.add_textbox(Inches(7.85), Inches(1.45), Inches(4.9), Inches(5.2))
        tf = tb.text_frame
        tf.word_wrap = True
        add_bullets(tf, bullets or [], size=15)
    if caption:
        tb = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = caption
        set_font(r, size=9.5, color=GRAY)
    add_notes(s, notes)
    add_footer(s)
    return s


def add_placeholder(slide, left, top, width, height, title, detail):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF7, 0xF9, 0xFC)
    box.line.color.rgb = ACCENT
    box.line.width = Pt(1.2)
    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, size=11, bold=True, color=ACCENT)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = detail
    set_font(r2, size=9, color=GRAY)
    return box


# ============================================================
# 1. 封面（项目名不变）
# ============================================================
s = prs.slides.add_slide(BLANK)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

tb = s.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.6), Inches(2.1))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "基于方向提示与 SAM 的"
set_font(r, size=29, bold=True, color=DARK)
p2 = tf.add_paragraph()
r2 = p2.add_run(); r2.text = "遥感影像阴影检测方法研究"
set_font(r2, size=29, bold=True, color=DARK)

tb2 = s.shapes.add_textbox(Inches(0.9), Inches(3.85), Inches(11.5), Inches(2.2))
tf2 = tb2.text_frame; tf2.word_wrap = True
for i, (t, sz, color) in enumerate([
    ("物理先验提示 + LoRA 微调 SAM 的全自动实现", 18, ACCENT),
    ("大学生创新创业训练计划项目 · 结题答辩", 16, GRAY),
    ("团队：ShadowHunter　|　2026 年 10 月", 15, GRAY),
]):
    pp = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    rr = pp.add_run(); rr.text = t
    set_font(rr, size=sz, color=color)
    pp.space_after = Pt(7)
add_notes(s, "各位老师好，我是 ShadowHunter 团队代表。今天汇报的项目是《基于方向提示与 SAM 的遥感影像阴影检测方法研究》。"
             "我会用 5 分钟讲清楚：我们做了什么、结果如何、边界在哪。")
add_footer(s)

# ============================================================
# 2. 问题与挑战
# ============================================================
image_slide(
    "阴影检测：下游任务的“隐形干扰源”",
    "results/stages/chicago22_sub9/01_original.png",
    bullets=[
        "阴影遮挡细节、扭曲光谱，影响分类、变化检测、目标识别",
        "传统阈值法：逐场景调参；水体、深色植被、深色路面易混淆",
        "全监督网络：像素级标注贵，跨场景迁移不稳",
        "SAM 边界分割强，但零样本并不懂“遥感阴影”",
    ],
    mode="left", image_h=4.7,
    caption="图：AISD 测试图 chicago22_sub9；建筑、树木与暗色地物同时存在。",
    notes="阴影检测的难点是类间相似、边界复杂、标注贵。传统方法靠阈值，深度方法靠标注，SAM 虽然强但零样本不懂阴影。"
          "本项目要解决的就是：如何用低成本把 SAM 变成可用的阴影检测器。约 20 秒。",
)

# ============================================================
# 3. 评估基座
# ============================================================
s = prs.slides.add_slide(BLANK)
title_bar(s, "先修评估基座：一个被忽视的 42/43 同源泄漏", "所有方法同划分重训、同脚本评估")
tb = s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(5.8), Inches(5.2))
tf = tb.text_frame; tf.word_wrap = True
add_bullets(tf, [
    "AISD 自带 Test51 与 Train412：43 个原始大图基名中 42 个重叠",
    "51 张测试图中 50 张来自训练场景 → IoU 系统性偏高",
    "自研基线 U-Net / SegNet 还存在图像-掩膜独立翻转的标签错位 bug",
    "按原始大图基名重划 68/16/21（262/58/92），三集不重叠",
], size=15)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.95))
box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = ACCENT; box.line.width = Pt(1)
tf = box.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.TOP
tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "为什么这是加分项"
set_font(r, size=16, bold=True, color=ACCENT)
for t in ["不是只报一个高分，而是先证明分数可信", "所有深度方法在同一划分上重训", "逐图/场景配对，测试集不参与选模型", "最终口径：92 张测试图 / 21 个独立场景"]:
    pp = tf.add_paragraph()
    rr = pp.add_run(); rr.text = "• " + t
    set_font(rr, size=13, color=DARK)
    pp.space_after = Pt(7)
add_notes(s, "这一页讲评估基座。AISD 自带的 Test51 和训练集 42/43 个基名重叠，51 张里 50 张同源，指标会虚高。"
             "我们按原始大图重划 68/16/21，所有方法同划分重训；还修了 U-Net/SegNet 的翻转 bug。这是后面所有数字可信的前提。约 25 秒。")
add_footer(s)

# ============================================================
# 4. 核心思路 + 方法总览
# ============================================================
s = prs.slides.add_slide(BLANK)
title_bar(s, "一条全自动链路：物理先验找候选，SAM 修边界", "无需人工点选，也不逐场景调阈值")
pic = s.shapes.add_picture(str(BASE / "results/figures/pipeline_trim.png"),
                           Inches(0.6), Inches(1.4), width=Inches(12.13))
if pic.height > Inches(3.6):
    scale = Inches(3.6) / pic.height
    pic.height = Inches(3.6); pic.width = int(pic.width * scale)
tb = s.shapes.add_textbox(Inches(0.6), Inches(4.55), Inches(12.13), Inches(2.2))
tf = tb.text_frame; tf.word_wrap = True
add_bullets(tf, [
    "① CLAHE + Otsu + 形态学：高召回候选阴影区",
    "② 连通域框提示：提示只来自输入影像，不读取 GT",
    "③ LoRA 微调 SAM：只训练 4.51M（约 4.8%）参数",
    "④ 掩膜矢量化：GeoJSON / Shapefile，保留孔洞与地理参考",
], size=15)
add_notes(s, "整体流程四步：物理粗提取得到高召回候选；连通域自动生成框提示，不读 GT；LoRA 微调 SAM 修边界；最后矢量化。"
             "全程不需要人工点选。约 25 秒。")
add_footer(s)

# ============================================================
# 5. 物理粗提取
# ============================================================
image_slide(
    "物理粗提取：先要“找得全”，再让 SAM “切得准”",
    "results/figures/stage_coarse.png",
    bullets=[
        "CLAHE 增强暗部细节",
        "反相 + Otsu 分离暗区",
        "形态学开运算去细小噪声",
        "连通域面积 ≥300 px 才保留",
        "目标是高召回候选，不是最终边界",
        "参数在训练/验证阶段固定，测试不调",
    ],
    mode="left", image_h=5.2,
    caption="图：原图 → 灰度 → CLAHE → 反相 Otsu → 形态学 → 框提示。",
    notes="前六步都是传统图像处理：CLAHE 增强暗部，反相 Otsu 分离暗区，形态学去噪，最后取连通域包围框。"
          "粗提取的目标是高召回，边界交给 SAM。约 20 秒。",
)

# ============================================================
# 6. 框提示 + SAM 精修
# ============================================================
image_slide(
    "从框提示到 SAM 精修：全自动、不读 GT",
    "results/figures/stage_sam.png",
    bullets=[
        "连通域轴对齐包围框 → SAM 框提示",
        "逐框预测后取并集，得到精细边界",
        "最终叠加：绿色为预测阴影",
        "误差图：TP 绿 / FP 红 / FN 蓝",
        "边界比粗提取干净；暗色地物仍会过检",
    ],
    mode="left", image_h=5.1,
    caption="图：框提示 → 逐框 SAM 并集 → 最终叠加 → 误差图。",
    notes="框提示的关键是“不读 GT”。我们只用输入影像做增强粗提取，再取连通域包围框；SAM 逐框分割后取并集。"
          "边界比粗提取干净很多，但植被、暗色地物仍会被并进来。约 15 秒。",
)

# ============================================================
# 7. LoRA
# ============================================================
image_slide(
    "LoRA 微调：4.8% 参数撬动 ViT-B",
    "results/figures/param_efficiency.png",
    bullets=[
        "编码器冻结，注意力层 qkv/proj 挂 rank=8 LoRA",
        "可训练：LoRA 0.44M + prompt encoder + mask decoder = 4.51M",
        "损失：Focal + Dice；点/框双提示；水平翻转；AMP",
        "全部训练在 RTX 4060 Laptop（8GB）完成",
        "低秩适配降低小样本过拟合风险",
    ],
    mode="left", image_h=4.7,
    caption="图：总参数 93.7M，可训练 4.51M（约 4.8%）。",
    notes="LoRA 是核心工程点。全量微调在 8GB 卡上不现实，所以我们冻结编码器，只训练低秩矩阵、提示编码器和解码器。"
          "这样既省显存，也降低在 262 张训练图上的过拟合风险。约 20 秒。",
)

# ============================================================
# 8. 实验设置
# ============================================================
table_slide(
    "实验设置：主口径只有一个",
    [
        ["环节", "设置"],
        ["数据", "AISD（无泄漏 68/16/21）+ RSISD 跨数据集"],
        ["主口径", "262 / 58 / 92 张；92 张测试图 / 21 个独立场景"],
        ["指标", "IoU / BER / F1 / OA / P / R；macro；best-val 选模型"],
        ["基线", "U-Net / SegNet / Mask R-CNN / 冻结 SAM / 零样本 / 传统方法"],
    ],
    widths=[2.2, 9.0], best_row=None, row_h=0.68, top=1.55, size=14,
    sub="测试集只用于最终报告，不参与选模型",
    note=["AISD 自带划分的数字只作历史对齐，禁与主表横比。", "主结果报 3 seed 均值；逐图/场景配对只作描述性统计。"],
    notes="实验设置这一页最重要的是“主口径只有一个”：无泄漏 68/16/21，92 张测试图、21 个独立场景。"
          "所有方法同划分重训，测试集不参与选模型。约 15 秒。",
)

# ============================================================
# 9. 主结果表
# ============================================================
table_slide(
    "主结果：同划分、同脚本，数值上高于三个全监督基线",
    [
        ["方法", "IoU", "BER", "F1", "种子"],
        ["LoRA 微调 SAM（本项目）", "0.7990 ± 0.0101", "0.0612 ± 0.0027", "0.8869", "3 seed"],
        ["全监督 U-Net", "0.7547 ± 0.0028", "0.0860 ± 0.0091", "0.8567", "3 seed"],
        ["全监督 SegNet", "0.7382 ± 0.0044", "0.1003 ± 0.0020", "0.8448", "3 seed"],
        ["全监督 Mask R-CNN", "0.7392 ± 0.0039", "0.0955 ± 0.0042", "0.8486", "3 seed"],
    ],
    widths=[3.6, 2.5, 2.4, 1.4, 1.2], best_row=1, row_h=0.62, top=1.5, size=13,
    sub="92 张 / 21 独立场景；best-val；macro；3 seed",
    note=[
        "配对检验：LoRA 在 87/92 张图上优于 Mask R-CNN；21/21 个独立场景 BER 更低（描述性统计，未做多重比较校正）。",
        "表述纪律：只说“数值上高于”，不宣称显著超越 / SOTA。",
    ],
    notes="主结果：IoU 0.799、BER 0.061、F1 0.887。三个基线分别是 0.755、0.738、0.739。"
          "3-seed 配对里，87/92 张图优于 Mask R-CNN，21/21 个场景 BER 更低。我们只说数值上高于。约 25 秒。",
)

# ============================================================
# 10. 主结果图 + 配对
# ============================================================
s = prs.slides.add_slide(BLANK)
title_bar(s, "3-seed 主结果：数值优势在 3 个 seed 上一致")
pic = s.shapes.add_picture(str(BASE / "results/figures/main_result_clean.png"),
                           Inches(0.6), Inches(1.35), width=Inches(12.13))
if pic.height > Inches(4.95):
    scale = Inches(4.95) / pic.height
    pic.height = Inches(4.95)
    pic.width = int(pic.width * scale)
pic.left = Inches((13.333 - pic.width / 914400) / 2)
tb = s.shapes.add_textbox(Inches(0.6), Inches(6.4), Inches(12.13), Inches(0.6))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "配对：LoRA 87/92 图胜 Mask R-CNN；21/21 场景 BER 更低（描述性统计，未做多重比较校正）。",
    "效应量随 pooled / macro / scene 口径变化；主结果只作数值优势表述。",
], size=12)
add_notes(s, "这一页看 3-seed 主结果：LoRA 的 IoU 0.799、BER 0.061，三个基线分别是 0.755、0.738、0.739。"
             "配对里 87/92 张图胜 Mask R-CNN，21/21 个场景 BER 更低。错误棒是 3-seed 样本标准差；"
             "效应量随口径变化，所以我们只说“数值上高于”。约 20 秒。")
add_footer(s)

# ============================================================
# 11. 消融
# ============================================================
table_slide(
    "消融：零样本 → 冻结编码器 → LoRA 的受控递进",
    [
        ["配置", "IoU", "BER", "F1", "说明"],
        ["零样本 SAM（同框提示）", "0.3480", "0.2923", "0.4997", "只差“是否微调”"],
        ["冻结编码器 + 训练提示/解码", "0.5621", "0.1327", "0.7105", "单 seed"],
        ["LoRA 微调 SAM（本项目）", "0.7990", "0.0612", "0.8869", "3 seed 均值"],
    ],
    widths=[3.8, 1.3, 1.3, 1.3, 2.2], best_row=3, row_h=0.66, top=1.55, size=14,
    sub="同一提示链路、同一划分、同一评估脚本",
    note=[
        "零样本 → LoRA：逐图配对 +0.4399 IoU（seed-0），92/92 张全部变好；3-seed 平均 +0.4510。",
        "冻结 → LoRA：LoRA 在 91/92 张更高。可训练参数包含 LoRA + prompt encoder + mask decoder。",
    ],
    notes="消融回答“增益从哪来”：零样本 0.348，冻结编码器 0.562，LoRA 0.799。"
          "零样本到 LoRA 是 92 张全部变好，冻结到 LoRA 是 91/92 张更高。关键不是换提示，而是让编码器适配遥感域。约 20 秒。",
)

# ============================================================
# 12. 跨数据集与误差
# ============================================================
table_slide(
    "跨数据集与误差：哪里有提升，哪里仍会错",
    [
        ["项目", "结果", "必须说明"],
        ["RSISD test200 零样本迁移", "IoU 0.7842（n=200, seed-0）", "任务较易：Otsu 即 0.5711；禁读作“泛化好”"],
        ["误检 / 漏检", "FPR 0.0405 / FNR 0.0882", "漏检约为误检 2.2 倍，主误差在召回侧"],
        ["小连通域召回", "<256 px：0.538", "但按漏检像素量，>=4096 px 贡献 39.0%"],
        ["边界带 F1", "r=1/3/5：0.745 / 0.819 / 0.849", "边界是难点，但不是唯一主项"],
    ],
    widths=[2.8, 3.8, 5.0], best_row=None, row_h=0.7, top=1.55, size=12,
    sub="均为无泄漏口径；本页数值为 seed-0 主方法；不同口径不可混读",
    notes="跨数据集 IoU 0.784，但这个任务本身容易，Otsu 都能到 0.571，所以不能读成泛化好。"
          "误差方面，漏检是误检的 2.2 倍；小阴影召回低，但按漏检像素量，大连通域贡献更多。约 15 秒。",
)

# ============================================================
# 13. 单图多方法
# ============================================================
image_slide(
    "同一张 AISD 影像：6 个关键方法的边界行为",
    "results/figures/compare_aisd_key.png",
    bullets=[
        "LoRA：IoU 0.676 / BER 0.081",
        "SegNet：IoU 0.677 / BER 0.123",
        "冻结编码器：0.345 / 0.167",
        "零样本：0.107 / 0.479",
        "U-Net / Mask R-CNN：IoU 0.695，略高于 LoRA",
        "n=1、seed-0；不排序、不标 best",
    ],
    mode="left", image_h=5.0,
    caption="图：chicago22_sub9（AISD 测试集）；单图数值只作陈列，不构成方法优劣结论。",
    notes="这是同一张 AISD 图上的 6 个关键方法。LoRA 的 BER 最低，但 IoU 不是最高。我们不做排序，因为 n=1。"
          "它主要用来看失败模式：植被和暗色地物仍会被并进阴影。约 15 秒。",
)

# ============================================================
# 14. 软件原型
# ============================================================
image_slide(
    "从掩膜到矢量：一个可运行的软件原型",
    "results/stages/chicago22_sub9/09_final_overlay.png",
    bullets=[
        "输入影像 → 阴影多边形（GeoJSON / SHP，含孔洞）",
        "42 个多边形；矢量↔栅格面积差 0.0107%",
        "栅格化 IoU 0.8438，对照管线 0.8453",
        "支持地理参考 CRS；无参考时输出像素坐标",
        "端到端验证覆盖 GeoJSON / Shapefile / 合成 CRS",
        "默认 seed-0 权重；验证限单图 + 合成 CRS",
    ],
    mode="left", image_h=4.7,
    caption="图：最终叠加示例；绿色为预测阴影。验证日志见 logs/verify_shadow_extract_*.log。",
    notes="工程交付方面，shadow_extract.py 输入影像输出 GeoJSON 或 SHP，保留孔洞和地理参考。"
          "验证了三条分支：GeoJSON、Shapefile、合成 CRS；面积差 0.0107%。目前只测了单张图和合成 CRS。约 15 秒。",
)

# ============================================================
# 15. 方向先验 + S-EO
# ============================================================
s = prs.slides.add_slide(BLANK)
title_bar(s, "方向先验：把问题做成可复核的预注册诊断", "不写“方向可学”，也不写“方向无效”")
tb = s.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(6.2), Inches(4.3))
tf = tb.text_frame; tf.word_wrap = True
add_bullets(tf, [
    "四种注入：几何过滤 / 方向点提示 / 方向点训练 / 方向 Adapter",
    "配套：翻转自监督、方向回归头、DTICT、标签一致性探针",
    "旧标签诊断：掩膜长轴 vs 真实太阳角，中位差 33.97°（逐连通域）",
    "S-EO 预注册试点：OMA 42 日期、JAX 4 AOI × 25 日期",
    "结论：预注册试点未达到阳性判据；方向模块未进入主链",
], size=14)
pic = s.shapes.add_picture(str(BASE / "results/seo_direction_pilot/direction_diagnostic.png"),
                           Inches(6.95), Inches(1.4), width=Inches(5.8))
if pic.height > Inches(5.35):
    scale = Inches(5.35) / pic.height
    pic.height = Inches(5.35)
    pic.width = int(pic.width * scale)
add_notes(s, "方向模块我们做了四种注入和一套探针。旧标签和真实太阳角中位差 33.97°，说明旧标签混入建筑朝向。"
             "为了把问题做干净，我们又用 S-EO 的真实太阳角做了预注册试点：OMA 同 AOI 点估计 2.99°，低于常数 5.75° 和高程 oracle 3.96°，"
             "但日期块 p=0.186 没过主门槛；JAX 和双向跨域都没有阳性。所以结论是：预注册试点未达到阳性判据，方向模块没有进入主链。"
             "这是一个可复核的负结果，不是失败。约 30 秒。")
add_footer(s)

# ============================================================
# 16. 结论与边界
# ============================================================
content_slide(
    "结论：一条可运行、主结果可复核的全自动链路",
    [
        "主结果：无泄漏 3-seed IoU 0.7990 / BER 0.0612；仅训练 4.8% 参数；8GB GPU 完成",
        "方法学：发现并修复同源泄漏与基线增强 bug；重建 68/16/21 可复核评估基座",
        "完整对照：零样本 → 冻结 → LoRA；零样本→LoRA 92/92 张变好",
        "工程：输入影像 → 阴影矢量；GeoJSON / SHP；端到端验证通过",
        "方向先验：四种注入 + S-EO 预注册试点；未观察到稳健跨场景信号；未进入主链",
        "边界：21 独立场景；主方法 BER 6.12%；未复现 SOTA；软件验证限单图 + 合成 CRS",
    ],
    sub="后续：跨传感器验证、暗色地物判别、工具包与论文/软著",
    notes="最后总结：我们交付了一条可运行、主结果可复核的链路，主结果 0.799/0.061，只训 4.8% 参数，8GB GPU 完成；"
          "同时修复了评估基座的两个问题。方向模块经过预注册试点，没有观察到稳健信号，所以没有进入主链。"
          "边界也很清楚：21 个场景、BER 6.12%、没有 SOTA、软件验证有限。约 20 秒。",
)
s = prs.slides[-1]
banner = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.25),
                            Inches(12.13), Inches(0.5))
banner.fill.solid()
banner.fill.fore_color.rgb = ACCENT
banner.line.fill.background()
btf = banner.text_frame
btf.vertical_anchor = MSO_ANCHOR.MIDDLE
bp = btf.paragraphs[0]
bp.alignment = PP_ALIGN.CENTER
br = bp.add_run()
br.text = "可复核：results/ 逐图 CSV · logs/ 训练日志 · 一键复现命令"
set_font(br, size=13, bold=True, color=WHITE)

# ============================================================
# 17. 致谢
# ============================================================
s = prs.slides.add_slide(BLANK)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.5))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "谢谢各位老师，请批评指正！"
set_font(r, size=34, bold=True, color=DARK)
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(4.15), Inches(11.3), Inches(0.8))
tf2 = tb2.text_frame
p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run(); r2.text = "代码、结果 CSV、训练日志与复现命令均可复核"
set_font(r2, size=14, color=GRAY)
add_notes(s, "我的汇报到此结束，谢谢各位老师，请批评指正。")
add_footer(s)

out = BASE / os.environ.get("PPT_OUT", "结题答辩PPT_最终版.pptx")
prs.save(str(out))
print(f"PPT 已生成：{out}  共 {len(prs.slides)} 页")
