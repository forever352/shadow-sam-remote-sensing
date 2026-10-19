# -*- coding: utf-8 -*-
"""由 结题报告.md 生成排版好的 Word 版结题报告。

流程：pandoc（标题/表格/图片/公式） -> python-docx（中文字体、版式、表格、图片宽度）
-> Word COM（更新目录域、导出 PDF 预览）。
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "结题报告.md"
OUT_DIR = Path(__file__).resolve().parent
REF = OUT_DIR / "_reference.docx"
TMP_MD = OUT_DIR / "_report_for_docx.md"
OUT = ROOT / "结题报告.docx"
PREVIEW_DIR = OUT_DIR / "preview"
PREVIEW_PDF = PREVIEW_DIR / "结题报告_preview.pdf"

BODY_CN = "宋体"
BODY_EN = "Times New Roman"
HEAD_CN = "黑体"
CODE_FONT = "Consolas"
BODY_SIZE = 12
TABLE_SIZE = 9
MAX_IMAGE_WIDTH = Cm(15.5)


def set_font_xml(rPr, cn: str, en: str) -> None:
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)


def style_font(style, cn: str, en: str, size: float | None = None,
               bold: bool | None = None) -> None:
    style.font.name = en
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    set_font_xml(style.element.get_or_add_rPr(), cn, en)


def run_font(run, cn: str, en: str, size: float | None = None,
             bold: bool | None = None) -> None:
    run.font.name = en
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    set_font_xml(run._element.get_or_add_rPr(), cn, en)


def build_reference() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)

    styles = doc.styles
    normal = styles["Normal"]
    # Normal 用小字号/紧凑行距，供目录域条目继承；正文段落在后处理里显式覆盖为 12pt/1.5。
    style_font(normal, BODY_CN, BODY_EN, 9.5)
    normal.paragraph_format.line_spacing = 0.95
    normal.paragraph_format.space_after = Pt(0)

    specs = {
        "Title": (HEAD_CN, 22, True),
        "Subtitle": (HEAD_CN, 14, False),
        "Heading 1": (HEAD_CN, 16, True),
        "Heading 2": (HEAD_CN, 14, True),
        "Heading 3": (HEAD_CN, 12, True),
        "Heading 4": (HEAD_CN, 12, True),
    }
    for name, (cn, size, bold) in specs.items():
        if name not in styles:
            continue
        st = styles[name]
        style_font(st, cn, cn, size, bold)
        st.font.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.keep_with_next = True
        if name == "Title":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(0)
            pf.space_after = Pt(12)
        elif name == "Subtitle":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after = Pt(18)
        else:
            pf.space_before = Pt(12)
            pf.space_after = Pt(6)

    for name in ("Block Text", "Quote"):
        if name in styles:
            st = styles[name]
            style_font(st, BODY_CN, BODY_EN, 10.5)
            st.paragraph_format.left_indent = Cm(0.8)
            st.paragraph_format.right_indent = Cm(0.4)
            st.paragraph_format.space_before = Pt(3)
            st.paragraph_format.space_after = Pt(3)
            st.paragraph_format.line_spacing = 1.25

    for name in ("Source Code", "Verbatim Char"):
        if name in styles:
            st = styles[name]
            style_font(st, CODE_FONT, CODE_FONT, 8.5)

    for name in ("Image Caption", "Table Caption", "Caption"):
        if name in styles:
            st = styles[name]
            style_font(st, BODY_CN, BODY_EN, 9)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.space_before = Pt(3)
            st.paragraph_format.space_after = Pt(9)

    toc_specs = {
        "toc 1": (11, True, 0.0, 3),
        "toc 2": (10.5, False, 0.55, 2),
        "toc 3": (10, False, 1.1, 1),
        "TOC 1": (11, True, 0.0, 3),
        "TOC 2": (10.5, False, 0.55, 2),
        "TOC 3": (10, False, 1.1, 1),
    }
    for name, (size, bold, indent_cm, after) in toc_specs.items():
        if name not in styles:
            continue
        st = styles[name]
        style_font(st, BODY_CN, BODY_EN, size, bold)
        st.paragraph_format.line_spacing = 1.1
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.left_indent = Cm(indent_cm)

    doc.save(REF)


def run_pandoc() -> None:
    text = SRC.read_text(encoding="utf-8")
    # 副标题不进入目录：临时去掉 ## 前缀，后处理再设为居中副标题。
    text = text.replace(
        "## 大学生创新创业训练计划项目结题报告",
        "大学生创新创业训练计划项目结题报告",
        1,
    )
    TMP_MD.write_text(text, encoding="utf-8")
    cmd = [
        "pandoc",
        str(TMP_MD),
        "-f", "markdown+tex_math_dollars+pipe_tables",
        "-t", "docx",
        "--reference-doc", str(REF),
        "--resource-path", str(ROOT),
        "--toc", "--toc-depth=3",
        "-M", "toc-title=目录",
        "-o", str(OUT),
    ]
    subprocess.run(cmd, check=True, cwd=str(ROOT))


def para_has_image(p: Paragraph) -> bool:
    return bool(p._p.xpath(".//w:drawing"))


def is_list(p: Paragraph) -> bool:
    pPr = p._p.pPr
    return pPr is not None and pPr.numPr is not None


def apply_paragraph_format(doc: Document) -> None:
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        text = p.text.strip()

        if style == "Title":
            for r in p.runs:
                run_font(r, HEAD_CN, HEAD_CN, 22, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if style.startswith("Heading") and text.startswith("基于方向提示与 SAM 的遥感影像阴影检测方法研究"):
            try:
                p.style = doc.styles["Title"]
            except KeyError:
                pass
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(36)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.first_line_indent = Pt(0)
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            r1 = p.add_run("基于方向提示与 SAM 的")
            run_font(r1, HEAD_CN, HEAD_CN, 22, True)
            r1.add_break()
            r2 = p.add_run("遥感影像阴影检测方法研究")
            run_font(r2, HEAD_CN, HEAD_CN, 22, True)
            continue

        if text == "大学生创新创业训练计划项目结题报告" and style.startswith("Body"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(18)
            for r in p.runs:
                run_font(r, HEAD_CN, HEAD_CN, 14, False)
            continue

        if style.startswith("Heading"):
            for r in p.runs:
                run_font(r, HEAD_CN, HEAD_CN)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue

        if "Source Code" in style:
            for r in p.runs:
                run_font(r, CODE_FONT, CODE_FONT, 8.5)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F5F5F5")
            p._p.get_or_add_pPr().append(shd)
            continue

        if para_has_image(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            continue

        # 图/表题注：图 X ... 或 表 X ...
        if re.match(r"^(图|表)\s*\d+", text):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(9)
            p.paragraph_format.keep_with_next = text.startswith("表")
            for r in p.runs:
                run_font(r, BODY_CN, BODY_EN, 9)
            continue

        if "Block Text" in style or "Quote" in style:
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                run_font(r, BODY_CN, BODY_EN, 10.5)
            continue

        # 正文 / 列表
        if is_list(p):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.75)
        elif re.match(r"^\[\d+\]", text):
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
        else:
            p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            run_font(r, BODY_CN, BODY_EN, BODY_SIZE)


def apply_tables(doc: Document) -> None:
    for table in doc.tables:
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if len(table.columns) == 2:
            table.autofit = False
            table.allow_autofit = False
            table.columns[0].width = Cm(3.2)
            table.columns[1].width = Cm(12.2)
        for i, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            if i == 0:
                tbl_header = OxmlElement("w:tblHeader")
                tbl_header.set(qn("w:val"), "true")
                trPr.append(tbl_header)
            for j, cell in enumerate(row.cells):
                if len(table.columns) == 2:
                    cell.width = Cm(3.2) if j == 0 else Cm(12.2)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        run_font(r, BODY_CN, BODY_EN, TABLE_SIZE, bold=(i == 0))


def apply_images(doc: Document) -> None:
    for shape in doc.inline_shapes:
        width = int(shape.width)
        if width <= int(MAX_IMAGE_WIDTH):
            continue
        ratio = int(MAX_IMAGE_WIDTH) / width
        shape.width = int(MAX_IMAGE_WIDTH)
        shape.height = int(shape.height * ratio)


def add_field(paragraph: Paragraph, instr: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_end)


def apply_header_footer(doc: Document) -> None:
    for sec in doc.sections:
        header = sec.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("基于方向提示与 SAM 的遥感影像阴影检测方法研究")
        run_font(r, BODY_CN, BODY_EN, 9)
        r.font.color.rgb = RGBColor(128, 128, 128)

        footer = sec.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run("第 ")
        run_font(r1, BODY_CN, BODY_EN, 9)
        add_field(fp, " PAGE ")
        r2 = fp.add_run(" 页")
        run_font(r2, BODY_CN, BODY_EN, 9)


def add_page_breaks(doc: Document) -> None:
    for p in doc.paragraphs:
        text = p.text.strip()
        if p.style.name.startswith("Heading 2"):
            if text == "摘要" or re.match(r"^(1\s*引言|参考文献|附录 [AB])", text):
                p.paragraph_format.page_break_before = True


def finalize_with_word() -> bool:
    try:
        import win32com.client as win32
    except Exception as exc:  # pragma: no cover
        print(f"[warn] pywin32 unavailable: {exc}")
        return False

    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    word = None
    doc = None
    try:
        word = win32.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(OUT))
        try:
            doc.Fields.Update()
        except Exception:
            pass
        try:
            if doc.TablesOfContents.Count:
                doc.TablesOfContents(1).Update()
        except Exception:
            pass
        doc.Save()
        doc.ExportAsFixedFormat(
            OutputFileName=str(PREVIEW_PDF),
            ExportFormat=17,  # wdExportFormatPDF
            OpenAfterExport=False,
            OptimizeFor=0,
            Range=0,
            Item=0,
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,
            DocStructureTags=True,
        )
        return True
    except Exception as exc:  # pragma: no cover
        print(f"[warn] Word finalize failed: {exc}")
        return False
    finally:
        if doc is not None:
            try:
                doc.Close(SaveChanges=0)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"source not found: {SRC}")
    build_reference()
    run_pandoc()
    doc = Document(str(OUT))
    apply_paragraph_format(doc)
    apply_tables(doc)
    apply_images(doc)
    apply_header_footer(doc)
    add_page_breaks(doc)
    doc.save(str(OUT))

    ok = finalize_with_word()
    print(f"docx: {OUT}")
    if ok:
        print(f"pdf preview: {PREVIEW_PDF}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
